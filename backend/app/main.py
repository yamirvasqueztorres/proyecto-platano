import hashlib
import hmac
import io
import json
import secrets
from collections import Counter, defaultdict
from contextlib import asynccontextmanager
from datetime import date, datetime, time, timedelta, timezone
from typing import Annotated

import pandas as pd
from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session

from .auth import create_access_token, current_user, hash_password, require_permission, verify_password
from .config import get_settings
from .database import Base, engine, get_db
from .models import AuditLog, Carrier, CatalogItem, CorrectiveAction, LotEvent, NonConformity, Notification, NotificationRead, ProductionLot, QualityParameter, QualityRecord, RecordStatus, Role, RolePermission, Supplier, User, Validation, utcnow
from .module_catalog import MODULES, validate_quality_payload
from .permissions import DEFAULT_ROLE_PERMISSIONS, PERMISSIONS
from .pdf_reports import render_quality_pdf
from .schemas import CarrierCreate, CarrierOut, ImportResult, LoginRequest, ParameterUpdate, RecordCreate, RecordOut, SupplierCreate, SupplierOut, TokenOut, UserCreate, UserOut, UserUpdate, ValidationCreate
from .workflows import router as workflow_router

settings = get_settings()

INITIAL_USERS = [
    ("admin", "ana.salazar@tropical.pe", "Ana Salazar", Role.ADMINISTRADOR),
    ("coordinador", "carlos.medina@tropical.pe", "Carlos Medina", Role.COORDINADOR),
    ("inspector", "lucia.ramos@tropical.pe", "Lucía Ramos", Role.INSPECTOR),
    ("revisor", "marco.vega@tropical.pe", "Marco Vega", Role.REVISOR),
]
DEMO_SUPPLIERS = [
    ("AGR-014", "José Paredes Huamán", "Valle del Chira · Sullana", "Comercio justo"),
    ("AGR-009", "Rosa Flores Mendoza", "Querecotillo · Sullana", "Convencional"),
    ("AGR-021", "Luis Torres Castillo", "Marcavelica · Sullana", "Comercio justo"),
    ("AGR-006", "Elena Ríos León", "Salitral · Sullana", "Convencional"),
]
DEMO_CARRIERS = [
    ("Luis Peña", "DNI 43821975", "P3A-201", "987654321"),
    ("Transportes del Chira", "RUC 20604571892", "B8K-742", "973245618"),
]
DEFAULT_PARAMETERS = [
    ("brix_min", 7, "°Brix", "Límite inferior de sólidos solubles"),
    ("brix_max", 15, "°Brix", "Límite superior de sólidos solubles"),
    ("humedad_min", 55, "%", "Humedad mínima aceptada"),
    ("humedad_max", 75, "%", "Humedad máxima aceptada"),
    ("peso_primera_min", 151, "g", "Peso mínimo de primera calidad"),
    ("temperatura_pulpa_max", 18, "°C", "Temperatura máxima de pulpa durante embolsado"),
]

DEFAULT_CATALOGS = [
    ("producto", "PLV", "Plátano verde", "Materia prima principal"),
    ("producto", "PLP", "Plátano pelado", "Producto en proceso"),
    ("defecto", "PEQ", "Fruto pequeño", "Clasificación dimensional"),
    ("defecto", "CUR", "Fruto curvo", "Desviación de forma"),
    ("defecto", "PIN", "Fruto pintón", "Madurez fuera de especificación"),
    ("defecto", "MEC", "Daño mecánico", "Golpe o corte visible"),
    ("etapa", "REC", "Recepción", "Ingreso y control de materia prima"),
    ("etapa", "SEL", "Selección", "Clasificación de materia prima"),
    ("etapa", "PEL", "Pelado", "Retiro de cáscara"),
    ("etapa", "EMB", "Embolsado", "Acondicionamiento y empaque"),
    ("etapa", "PT", "Producto terminado", "Liberación de calidad"),
    ("etapa", "DES", "Despacho", "Salida de planta"),
    ("turno", "MAN", "Mañana", "06:00 a 14:00"),
    ("turno", "TAR", "Tarde", "14:00 a 22:00"),
]


def seed_database(db: Session) -> None:
    """Create required configuration; fictional business data requires explicit opt-in."""
    for username, email, name, role in INITIAL_USERS:
        if not db.scalar(select(User).where(User.username == username)):
            db.add(User(username=username, email=email, full_name=name, role=role, hashed_password=hash_password(settings.initial_password)))
    for key, value, unit, description in DEFAULT_PARAMETERS:
        if not db.scalar(select(QualityParameter).where(QualityParameter.key == key)):
            db.add(QualityParameter(key=key, value=value, unit=unit, description=description))
    for role, default_permissions in DEFAULT_ROLE_PERMISSIONS.items():
        for permission, _ in PERMISSIONS:
            if not db.scalar(select(RolePermission).where(RolePermission.role == role, RolePermission.permission == permission)):
                db.add(RolePermission(role=role, permission=permission, allowed=permission in default_permissions))
    for catalog_type, code, name, description in DEFAULT_CATALOGS:
        if not db.scalar(select(CatalogItem).where(CatalogItem.catalog_type == catalog_type, CatalogItem.code == code)):
            db.add(CatalogItem(catalog_type=catalog_type, code=code, name=name, description=description))
    db.commit()

    if not settings.seed_demo_data:
        return

    for code, name, origin, cert in DEMO_SUPPLIERS:
        if not db.scalar(select(Supplier).where(Supplier.code == code)):
            db.add(Supplier(code=code, name=name, origin=origin, certification_type=cert))
    for name, document, plate, phone in DEMO_CARRIERS:
        if not db.scalar(select(Carrier).where(Carrier.plate == plate)):
            db.add(Carrier(name=name, document=document, plate=plate, phone=phone))
    db.flush()
    admin = db.scalar(select(User).where(User.username == "admin"))
    supplier = db.scalar(select(Supplier).where(Supplier.code == "AGR-014"))
    if admin and not db.scalar(select(ProductionLot).where(ProductionLot.code == "L-0309-04")):
        lot = ProductionLot(code="L-0309-04", supplier_id=supplier.id if supplier else None, product="Plátano verde", certification_type="Comercio justo", harvest_date=date(2026, 9, 2), received_at=datetime(2026, 9, 3, 7, 35, tzinfo=timezone.utc), quantity_kg=1280, current_stage="Embolsado", status="Activo", observations="Lote demostrativo para recorrido de trazabilidad", created_by_id=admin.id)
        db.add(lot); db.flush()
        for stage, hour, description in [("Recepción", 7, "Ingreso y pesaje de materia prima"), ("Selección", 8, "Clasificación de primera y segunda"), ("Pelado", 9, "Control de calidad durante pelado"), ("Embolsado", 11, "Lote en acondicionamiento")]:
            db.add(LotEvent(lot_id=lot.id, stage=stage, event_type="Avance", description=description, occurred_at=datetime(2026, 9, 3, hour, 35, tzinfo=timezone.utc), created_by_id=admin.id))
    if admin and not db.scalar(select(NonConformity).where(NonConformity.code == "NC-2026-00001")):
        nc = NonConformity(code="NC-2026-00001", module_code="FOR-CCD-019", lot_code="L-0309-04", category="Fruto pequeño", severity="Mayor", description="Incidencia de frutos pequeños por encima del límite de control", quantity=42, status="En tratamiento", root_cause="Variación de calibre en la materia prima recibida", detected_at=datetime(2026, 9, 3, 10, 42, tzinfo=timezone.utc), detected_by_id=admin.id)
        db.add(nc); db.flush()
        db.add(CorrectiveAction(code="AC-2026-00001", nonconformity_id=nc.id, title="Reducir incidencia de calibre bajo", plan="Reforzar la segregación por calibre desde recepción y verificar la muestra de ingreso.", do_action="Capacitar al equipo de selección y colocar patrón visual de calibre.", responsible_id=admin.id, due_date=date.today() + timedelta(days=7), phase="Hacer", status="En curso", created_by_id=admin.id))
    db.commit()


@asynccontextmanager
async def lifespan(_: FastAPI):
    if settings.environment.lower() != "testing" and engine.url.get_backend_name() != "postgresql":
        raise RuntimeError("DATABASE_URL debe apuntar a PostgreSQL; SQLite solo está permitido con ENVIRONMENT=testing")
    if settings.environment.lower() == "production" and (len(settings.secret_key) < 32 or any(word in settings.secret_key.lower() for word in ("cambiar", "reemplace"))):
        raise RuntimeError("En producción SECRET_KEY debe ser aleatoria, segura y tener al menos 32 caracteres")
    with engine.begin() as connection:
        if connection.dialect.name == "postgresql":
            # Serialize first-time setup across processes; release on commit/error.
            connection.execute(text("SELECT pg_advisory_xact_lock(36020260904)"))
        Base.metadata.create_all(connection)
        with Session(bind=connection) as db:
            seed_database(db)
    yield


app = FastAPI(title=settings.app_name, version="2.0.0", description="API REST para el control digital de calidad y mejora continua del procesamiento de plátano verde.", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=settings.origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


def audit(db: Session, request: Request, user: User | None, action: str, entity: str, entity_id: str | None = None, details: dict | None = None) -> None:
    db.add(AuditLog(user_id=user.id if user else None, action=action, entity=entity, entity_id=entity_id, details=details or {}, ip_address=request.client.host if request.client else None))


def parameters(db: Session) -> dict[str, float]:
    return {item.key: item.value for item in db.scalars(select(QualityParameter)).all()}


def record_query(start: date | None = None, end: date | None = None, module_code: str | None = None, supplier_id: int | None = None, status: str | None = None):
    query = select(QualityRecord).order_by(QualityRecord.created_at.desc())
    if start: query = query.where(QualityRecord.created_at >= datetime.combine(start, time.min, tzinfo=timezone.utc))
    if end: query = query.where(QualityRecord.created_at < datetime.combine(end + timedelta(days=1), time.min, tzinfo=timezone.utc))
    if module_code: query = query.where(QualityRecord.module_code == module_code)
    if supplier_id: query = query.where(QualityRecord.supplier_id == supplier_id)
    if status:
        try:
            query = query.where(QualityRecord.status == RecordStatus(status))
        except ValueError:
            raise HTTPException(422, "Estado de registro no reconocido")
    return query


@app.get("/api/health", tags=["Sistema"])
def health(db: Session = Depends(get_db)):
    db.scalar(select(1))
    return {"status": "ok", "service": settings.app_name, "environment": settings.environment, "database": engine.url.get_backend_name(), "time": utcnow()}


@app.post("/api/auth/login", response_model=TokenOut, tags=["Autenticación"])
def login(body: LoginRequest, request: Request, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == body.username.lower()))
    if not user or not user.is_active or not verify_password(body.password, user.hashed_password):
        audit(db, request, user, "LOGIN_FAILED", "session", details={"username": body.username})
        db.commit()
        raise HTTPException(status_code=401, detail="Usuario o contraseña incorrectos")
    audit(db, request, user, "LOGIN", "session")
    db.commit()
    return TokenOut(access_token=create_access_token(user), user=user)


@app.get("/api/auth/me", response_model=UserOut, tags=["Autenticación"])
def me(user: User = Depends(current_user)):
    return user


@app.get("/api/modules", tags=["Módulos"])
def list_modules(_: User = Depends(current_user)):
    return [{"code": code, **definition} for code, definition in MODULES.items()]


@app.get("/api/users", response_model=list[UserOut], tags=["Usuarios"])
def list_users(db: Session = Depends(get_db), _: User = Depends(require_permission("users.manage"))):
    return db.scalars(select(User).order_by(User.full_name)).all()


@app.post("/api/users", response_model=UserOut, status_code=201, tags=["Usuarios"])
def create_user(body: UserCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("users.manage"))):
    if db.scalar(select(User).where(or_(User.username == body.username.lower(), User.email == body.email.lower()))):
        raise HTTPException(status_code=409, detail="El usuario o correo ya está registrado")
    user = User(username=body.username.lower(), email=body.email.lower(), full_name=body.full_name, role=body.role, hashed_password=hash_password(body.password))
    db.add(user); db.flush(); audit(db, request, actor, "CREATE", "user", str(user.id), {"role": body.role.value}); db.commit(); db.refresh(user)
    return user


@app.patch("/api/users/{user_id}", response_model=UserOut, tags=["Usuarios"])
def update_user(user_id: int, body: UserUpdate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("users.manage"))):
    user = db.get(User, user_id)
    if not user: raise HTTPException(404, "Usuario no encontrado")
    changes = body.model_dump(exclude_none=True)
    if "email" in changes:
        changes["email"] = str(changes["email"]).lower()
        duplicate = db.scalar(select(User).where(User.email == changes["email"], User.id != user.id))
        if duplicate: raise HTTPException(409, "El correo ya está registrado")
    if "password" in changes: changes["hashed_password"] = hash_password(changes.pop("password"))
    for key, value in changes.items(): setattr(user, key, value)
    audit(db, request, actor, "UPDATE", "user", str(user.id), {"fields": list(changes)})
    db.commit(); db.refresh(user); return user


@app.patch("/api/users/{user_id}/status", response_model=UserOut, tags=["Usuarios"])
def set_user_status(user_id: int, active: bool, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("users.manage"))):
    user = db.get(User, user_id)
    if not user: raise HTTPException(404, "Usuario no encontrado")
    if user.id == actor.id and not active: raise HTTPException(400, "No puede desactivar su propia cuenta")
    user.is_active = active; audit(db, request, actor, "UPDATE_STATUS", "user", str(user.id), {"active": active}); db.commit(); db.refresh(user)
    return user


@app.get("/api/suppliers", response_model=list[SupplierOut], tags=["Proveedores"])
def list_suppliers(q: str | None = None, db: Session = Depends(get_db), _: User = Depends(current_user)):
    query = select(Supplier).where(Supplier.is_active.is_(True)).order_by(Supplier.name)
    if q: query = query.where(or_(Supplier.name.ilike(f"%{q}%"), Supplier.code.ilike(f"%{q}%"), Supplier.origin.ilike(f"%{q}%")))
    return db.scalars(query).all()


@app.post("/api/suppliers", response_model=SupplierOut, status_code=201, tags=["Proveedores"])
def create_supplier(body: SupplierCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    if db.scalar(select(Supplier).where(Supplier.code == body.code.upper())): raise HTTPException(409, "El código del proveedor ya existe")
    values = body.model_dump()
    values["code"] = body.code.upper()
    item = Supplier(**values); db.add(item); db.flush(); audit(db, request, actor, "CREATE", "supplier", str(item.id)); db.commit(); db.refresh(item)
    return item


@app.patch("/api/suppliers/{supplier_id}", response_model=SupplierOut, tags=["Proveedores"])
def update_supplier(supplier_id: int, body: SupplierCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    item = db.get(Supplier, supplier_id)
    if not item: raise HTTPException(404, "Proveedor no encontrado")
    changes = body.model_dump(); changes["code"] = changes["code"].upper()
    duplicate = db.scalar(select(Supplier).where(Supplier.code == changes["code"], Supplier.id != item.id))
    if duplicate: raise HTTPException(409, "El código del proveedor ya existe")
    for key, value in changes.items(): setattr(item, key, value)
    audit(db, request, actor, "UPDATE", "supplier", str(item.id)); db.commit(); db.refresh(item); return item


@app.delete("/api/suppliers/{supplier_id}", status_code=204, tags=["Proveedores"])
def deactivate_supplier(supplier_id: int, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    item = db.get(Supplier, supplier_id)
    if not item: raise HTTPException(404, "Proveedor no encontrado")
    item.is_active = False; audit(db, request, actor, "DEACTIVATE", "supplier", str(item.id)); db.commit()


@app.get("/api/carriers", response_model=list[CarrierOut], tags=["Transportistas"])
def list_carriers(q: str | None = None, db: Session = Depends(get_db), _: User = Depends(current_user)):
    query = select(Carrier).where(Carrier.is_active.is_(True)).order_by(Carrier.name)
    if q: query = query.where(or_(Carrier.name.ilike(f"%{q}%"), Carrier.plate.ilike(f"%{q}%")))
    return db.scalars(query).all()


@app.post("/api/carriers", response_model=CarrierOut, status_code=201, tags=["Transportistas"])
def create_carrier(body: CarrierCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    item = Carrier(**body.model_dump()); db.add(item); db.flush(); audit(db, request, actor, "CREATE", "carrier", str(item.id)); db.commit(); db.refresh(item); return item


@app.patch("/api/carriers/{carrier_id}", response_model=CarrierOut, tags=["Transportistas"])
def update_carrier(carrier_id: int, body: CarrierCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    item = db.get(Carrier, carrier_id)
    if not item: raise HTTPException(404, "Transportista no encontrado")
    for key, value in body.model_dump().items(): setattr(item, key, value)
    audit(db, request, actor, "UPDATE", "carrier", str(item.id)); db.commit(); db.refresh(item); return item


@app.delete("/api/carriers/{carrier_id}", status_code=204, tags=["Transportistas"])
def deactivate_carrier(carrier_id: int, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("catalogs.manage"))):
    item = db.get(Carrier, carrier_id)
    if not item: raise HTTPException(404, "Transportista no encontrado")
    item.is_active = False; audit(db, request, actor, "DEACTIVATE", "carrier", str(item.id)); db.commit()


@app.post("/api/records", response_model=RecordOut, status_code=201, tags=["Registros"])
def create_record(body: RecordCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("records.create"))):
    try:
        conformity, alerts, computed = validate_quality_payload(body.module_code, body.payload, parameters(db))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    if body.supplier_id is not None:
        supplier = db.get(Supplier, body.supplier_id)
        if not supplier or not supplier.is_active:
            raise HTTPException(422, "Proveedor no encontrado o inactivo")
    if body.carrier_id is not None:
        carrier = db.get(Carrier, body.carrier_id)
        if not carrier or not carrier.is_active:
            raise HTTPException(422, "Transportista no encontrado o inactivo")
    if body.lot is not None:
        body.lot = body.lot.strip().upper() or None
    now = utcnow()
    if body.started_at > now + timedelta(minutes=2): raise HTTPException(422, "La hora de inicio no puede estar en el futuro")
    temp_code = f"TMP-{secrets.token_hex(8)}"
    record = QualityRecord(record_code=temp_code, module_code=body.module_code, supplier_id=body.supplier_id, carrier_id=body.carrier_id, lot=body.lot, certification_type=body.certification_type, payload=computed, conformity_percent=conformity, alerts=alerts, started_at=body.started_at, completed_at=now, duration_seconds=max(0, int((now - body.started_at).total_seconds())), created_by_id=actor.id)
    db.add(record); db.flush(); record.record_code = f"REG-{now.year}-{record.id:06d}"
    if body.lot:
        lot_code = body.lot.strip().upper()
        lot_item = db.scalar(select(ProductionLot).where(ProductionLot.code == lot_code))
        if not lot_item:
            lot_item = ProductionLot(code=lot_code, supplier_id=body.supplier_id, product=str(computed.get("producto") or "Plátano verde"), certification_type=body.certification_type, current_stage=MODULES[body.module_code]["name"], status="Activo", created_by_id=actor.id)
            db.add(lot_item); db.flush()
        db.add(LotEvent(lot_id=lot_item.id, stage=MODULES[body.module_code]["name"], event_type="Registro de calidad", description=f"Se vinculó {record.record_code}", record_code=record.record_code, created_by_id=actor.id))
        lot_item.current_stage = MODULES[body.module_code]["name"]
    if alerts:
        nc = NonConformity(code=f"TMP-{secrets.token_hex(6)}", record_id=record.id, module_code=body.module_code, lot_code=body.lot.strip().upper() if body.lot else None, category="Desviación de calidad", severity="Mayor", description="; ".join(alerts), status="Abierta", detected_by_id=actor.id)
        db.add(nc); db.flush(); nc.code = f"NC-{now.year}-{nc.id:05d}"
    db.add(Notification(title="Registro pendiente de validación", message=f"{record.record_code} · {MODULES[body.module_code]['name']}", level="warning"))
    audit(db, request, actor, "CREATE", "quality_record", str(record.id), {"code": record.record_code, "module": body.module_code, "alerts": alerts})
    db.commit(); db.refresh(record)
    return record


@app.get("/api/records", response_model=list[RecordOut], tags=["Registros"])
def list_records(start: date | None = None, end: date | None = None, module_code: str | None = None, supplier_id: int | None = None, status: str | None = None, limit: Annotated[int, Query(ge=1, le=500)] = 100, db: Session = Depends(get_db), _: User = Depends(require_permission("records.view"))):
    return db.scalars(record_query(start, end, module_code, supplier_id, status).limit(limit)).all()


@app.get("/api/records/{record_id}", response_model=RecordOut, tags=["Registros"])
def get_record(record_id: int, db: Session = Depends(get_db), _: User = Depends(require_permission("records.view"))):
    record = db.get(QualityRecord, record_id)
    if not record: raise HTTPException(404, "Registro no encontrado")
    return record


@app.post("/api/records/{record_id}/validate", tags=["Validaciones"])
def validate_record(record_id: int, body: ValidationCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("records.validate"))):
    record = db.get(QualityRecord, record_id)
    if not record: raise HTTPException(404, "Registro no encontrado")
    if record.status == RecordStatus.APROBADO: raise HTTPException(409, "El registro ya se encuentra aprobado")
    if body.decision != "Aprobado" and not (body.observations or "").strip():
        raise HTTPException(422, "Debe registrar el motivo de la observación o rechazo")
    now = utcnow(); code = f"VAL-{now:%Y%m%d}-{secrets.token_hex(3).upper()}"
    signature_source = f"{record.record_code}|{actor.id}|{body.decision}|{now.isoformat()}".encode()
    signature = hmac.new(settings.secret_key.encode(), signature_source, hashlib.sha256).hexdigest()
    validation = Validation(record_id=record.id, decision=body.decision, observations=body.observations, validation_code=code, signature_hash=signature, validated_by_id=actor.id, validated_at=now)
    record.status = RecordStatus(body.decision); db.add(validation)
    db.add(Notification(user_id=record.created_by_id, title=f"Registro {body.decision.lower()}", message=f"{record.record_code} fue revisado por {actor.full_name}", level="success" if body.decision == "Aprobado" else "warning"))
    audit(db, request, actor, "VALIDATE", "quality_record", str(record.id), {"decision": body.decision, "validation_code": code}); db.commit()
    return {"record_code": record.record_code, "status": record.status.value, "validation_code": code, "signature_hash": signature, "validated_at": now}


@app.get("/api/dashboard", tags=["Dashboard"])
def dashboard(start: date | None = None, end: date | None = None, db: Session = Depends(get_db), _: User = Depends(require_permission("dashboard.view"))):
    records = db.scalars(record_query(start, end)).all()
    valid = [r.conformity_percent for r in records if r.conformity_percent is not None]
    weights = [float(r.payload.get("peso_promedio")) for r in records if r.payload.get("peso_promedio") not in (None, "")]
    pending = sum(r.status in (RecordStatus.PENDIENTE, RecordStatus.OBSERVADO) for r in records)
    by_day: dict[str, list[float]] = defaultdict(list)
    defect_counter: Counter[str] = Counter()
    for record in records:
        if record.conformity_percent is not None: by_day[record.created_at.date().isoformat()].append(record.conformity_percent)
        for key, value in record.payload.items():
            if key.startswith(("defecto_", "nc_")):
                try: defect_counter[key] += float(value or 0)
                except (ValueError, TypeError): pass
    open_nc = db.scalar(select(func.count()).select_from(NonConformity).where(NonConformity.status != "Cerrada")) or 0
    overdue_actions = db.scalar(select(func.count()).select_from(CorrectiveAction).where(CorrectiveAction.status != "Completada", CorrectiveAction.due_date < date.today())) or 0
    active_lots = db.scalar(select(func.count()).select_from(ProductionLot).where(ProductionLot.status.in_(["Activo", "Retenido", "Liberado"]))) or 0
    return {
        "conformity_percent": round(sum(valid) / len(valid), 2) if valid else None,
        "average_weight_g": round(sum(weights) / len(weights), 2) if weights else None,
        "pending_validations": pending,
        "average_entry_seconds": round(sum(r.duration_seconds for r in records) / len(records), 1) if records else None,
        "records_total": len(records),
        "open_nonconformities": open_nc,
        "overdue_actions": overdue_actions,
        "active_lots": active_lots,
        "trend": [{"date": day, "conformity": round(sum(values)/len(values), 2)} for day, values in sorted(by_day.items())],
        "top_defects": [{"field": key, "count": value} for key, value in defect_counter.most_common(10)],
    }


@app.get("/api/notifications", tags=["Notificaciones"])
def notifications(db: Session = Depends(get_db), actor: User = Depends(require_permission("alerts.view"))):
    query = select(Notification).where(or_(Notification.user_id == actor.id, Notification.user_id.is_(None))).order_by(Notification.created_at.desc()).limit(50)
    reads = set(db.scalars(select(NotificationRead.notification_id).where(NotificationRead.user_id == actor.id)).all())
    return [{"id": item.id, "user_id": item.user_id, "title": item.title, "message": item.message, "level": item.level, "is_read": item.is_read if item.user_id is not None else item.id in reads, "created_at": item.created_at} for item in db.scalars(query).all()]


@app.patch("/api/notifications/{notification_id}/read", tags=["Notificaciones"])
def read_notification(notification_id: int, db: Session = Depends(get_db), actor: User = Depends(require_permission("alerts.view"))):
    item = db.get(Notification, notification_id)
    if not item or item.user_id not in (None, actor.id): raise HTTPException(404, "Notificación no encontrada")
    if item.user_id is None:
        if not db.scalar(select(NotificationRead).where(NotificationRead.notification_id == item.id, NotificationRead.user_id == actor.id)):
            db.add(NotificationRead(notification_id=item.id, user_id=actor.id))
    else:
        item.is_read = True
    db.commit(); return {"ok": True}


@app.get("/api/audit", tags=["Auditoría"])
def audit_log(limit: Annotated[int, Query(ge=1, le=1000)] = 200, db: Session = Depends(get_db), _: User = Depends(require_permission("audit.view"))):
    return db.scalars(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(limit)).all()


@app.get("/api/parameters", tags=["Configuración"])
def get_parameters(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return db.scalars(select(QualityParameter).order_by(QualityParameter.key)).all()


@app.patch("/api/parameters/{key}", tags=["Configuración"])
def update_parameter(key: str, body: ParameterUpdate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_permission("parameters.manage"))):
    item = db.scalar(select(QualityParameter).where(QualityParameter.key == key))
    if not item: raise HTTPException(404, "Parámetro no encontrado")
    previous = item.value; item.value = body.value; audit(db, request, actor, "UPDATE", "quality_parameter", key, {"before": previous, "after": body.value}); db.commit(); db.refresh(item)
    return item


@app.post("/api/import/records", response_model=ImportResult, tags=["Importación"])
async def import_records(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db), actor: User = Depends(require_permission("records.import"))):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls", ".csv")): raise HTTPException(415, "Use un archivo Excel o CSV")
    content = await file.read()
    if len(content) > 15 * 1024 * 1024: raise HTTPException(413, "El archivo supera el máximo de 15 MB")
    frame = pd.read_csv(io.BytesIO(content)) if file.filename.lower().endswith(".csv") else pd.read_excel(io.BytesIO(content))
    imported, errors = 0, []
    for index, row in frame.fillna("").iterrows():
        try:
            with db.begin_nested():
                module_code = str(row.get("module_code", "")).strip()
                raw_payload = row.get("payload_json", "{}")
                payload = json.loads(raw_payload) if isinstance(raw_payload, str) else dict(raw_payload)
                conformity, alerts, computed = validate_quality_payload(module_code, payload, parameters(db))
                now = utcnow(); temp = f"IMP-{secrets.token_hex(8)}"
                record = QualityRecord(record_code=temp, module_code=module_code, lot=str(row.get("lot") or "") or None, certification_type=str(row.get("certification_type") or "Convencional"), payload=computed, conformity_percent=conformity, alerts=alerts, started_at=now, completed_at=now, duration_seconds=0, created_by_id=actor.id)
                db.add(record); db.flush(); record.record_code = f"REG-{now.year}-{record.id:06d}"
            imported += 1
        except Exception as exc:
            errors.append({"row": int(index) + 2, "error": str(exc)})
    audit(db, request, actor, "IMPORT", "quality_record", details={"file": file.filename, "imported": imported, "rejected": len(errors)}); db.commit()
    return ImportResult(rows_received=len(frame), rows_imported=imported, rows_rejected=len(errors), errors=errors[:100])


def report_rows(db: Session, start: date | None, end: date | None, module_code: str | None, supplier_id: int | None) -> list[QualityRecord]:
    return db.scalars(record_query(start, end, module_code, supplier_id)).all()


@app.get("/api/reports/excel", tags=["Reportes"])
def excel_report(start: date | None = None, end: date | None = None, module_code: str | None = None, supplier_id: int | None = None, db: Session = Depends(get_db), _: User = Depends(require_permission("reports.export"))):
    rows = report_rows(db, start, end, module_code, supplier_id)
    book = Workbook(); sheet = book.active; sheet.title = "Registros de calidad"
    headers = ["Código", "Módulo", "Fecha", "Estado", "Proveedor", "Lote", "Certificación", "Conformidad %", "Duración (s)", "Alertas", "Datos del módulo"]
    sheet.append(headers)
    for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="14795F"); cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append([row.record_code, row.module_code, row.created_at.isoformat(), row.status.value, row.supplier.name if row.supplier else "", row.lot or "", row.certification_type, row.conformity_percent, row.duration_seconds, " | ".join(row.alerts), json.dumps(row.payload, ensure_ascii=False)])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for col, width in {"A":20,"B":18,"C":24,"D":14,"E":28,"F":16,"G":18,"H":16,"I":16,"J":45,"K":70}.items(): sheet.column_dimensions[col].width = width
    def add_sheet(title: str, headers: list[str], data: list[list], widths: list[int]) -> None:
        target = book.create_sheet(title)
        target.append(headers)
        for cell in target[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="14795F"); cell.alignment = Alignment(horizontal="center")
        for values in data:
            target.append(values)
        target.freeze_panes = "A2"; target.auto_filter.ref = target.dimensions
        for index, width in enumerate(widths, start=1):
            target.column_dimensions[chr(64 + index)].width = width
    nonconformities = db.scalars(select(NonConformity).order_by(NonConformity.detected_at.desc())).all()
    actions = db.scalars(select(CorrectiveAction).order_by(CorrectiveAction.due_date)).all()
    lots = db.scalars(select(ProductionLot).order_by(ProductionLot.created_at.desc())).all()
    add_sheet("No conformidades", ["Código","Fecha","Módulo","Lote","Categoría","Severidad","Estado","Descripción","Causa raíz","Disposición"], [[n.code,n.detected_at.isoformat(),n.module_code,n.lot_code,n.category,n.severity,n.status,n.description,n.root_cause,n.disposition] for n in nonconformities], [18,23,18,16,20,12,18,45,45,45])
    add_sheet("Acciones PHVA", ["Código","NC","Título","Responsable","Vencimiento","Fase","Estado","Eficacia %","Planificar","Hacer","Verificar","Actuar"], [[a.code,a.nonconformity.code if a.nonconformity else None,a.title,a.responsible.full_name,a.due_date.isoformat(),a.phase,a.status,a.effectiveness_percent,a.plan,a.do_action,a.check_result,a.act_standardization] for a in actions], [18,18,30,24,14,14,16,12,45,45,45,45])
    add_sheet("Lotes", ["Código","Proveedor","Producto","Certificación","Recepción","Cantidad kg","Etapa","Estado","Observaciones"], [[lot.code,lot.supplier.name if lot.supplier else None,lot.product,lot.certification_type,lot.received_at.isoformat(),lot.quantity_kg,lot.current_stage,lot.status,lot.observations] for lot in lots], [18,28,22,18,23,14,22,14,45])
    output = io.BytesIO(); book.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=reporte_calidad.xlsx"})


@app.get("/api/reports/pdf", tags=["Reportes"])
def pdf_report(start: date | None = None, end: date | None = None, module_code: str | None = None, supplier_id: int | None = None, db: Session = Depends(get_db), _: User = Depends(require_permission("reports.export"))):
    rows = report_rows(db, start, end, module_code, supplier_id)
    nonconformities = db.scalars(select(NonConformity).order_by(NonConformity.detected_at.desc())).all()
    actions = db.scalars(select(CorrectiveAction).order_by(CorrectiveAction.due_date)).all()
    content = render_quality_pdf(rows, nonconformities, actions, utcnow())
    return StreamingResponse(io.BytesIO(content), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=reporte_calidad.pdf"})


@app.get("/", response_class=HTMLResponse, include_in_schema=False)
def root():
    return "<h1>Calidad 360 API</h1><p>Documentación técnica: <a href='/docs'>/docs</a></p>"


app.include_router(workflow_router)
